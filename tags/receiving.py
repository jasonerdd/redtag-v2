"""Parse and apply receiving Excel sheets into Vehicle profiles."""

from __future__ import annotations

import datetime
import re
from dataclasses import asdict, dataclass, field
from typing import Any

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .models import Dealer, Vehicle, VehicleModel

HEADER_ALIASES = {
    'customer': ('customer', 'dealer'),
    'model': ('model',),
    'vehicle_type': ('vehicle type', 'vehicle_type', 'type'),
    'lot': ('lot no.', 'lot no', 'lot', 'lot number'),
    'chassis': ('chassis no', 'chassis no.', 'chassis', 'chassis number'),
    'engine_no': ('engine no', 'engine no.', 'engine'),
    'arrival_date': ('ckd arrival date', 'arrival date', 'ckd arrival', 'receiving date'),
    'colour': ('colour', 'color'),
    'bond_entry': ('bond entry (im7)', 'bond entry', 'im7'),
    'receiving_status': ('receiving status', 'status'),
}

TEMPLATE_HEADERS = [
    'CUSTOMER',
    'MODEL',
    'VEHICLE TYPE',
    'LOT NO.',
    'CHASSIS NO',
    'ENGINE NO',
    'CKD ARRIVAL DATE',
]


def normalize_chassis(value: Any) -> str:
    """Canonical chassis: last 7 characters of the cleaned value."""
    text = _as_text(value)
    if not text:
        return ''
    return text[-7:] if len(text) > 7 else text


def find_vehicle_by_chassis(chassis: str):
    """Match exact chassis, then a longer full VIN that ends with the short form."""
    raw = (chassis or '').strip()
    short = normalize_chassis(raw)
    if not short:
        return None
    vehicle = Vehicle.objects.filter(chassis_no=raw).select_related('model', 'dealer').first()
    if vehicle:
        return vehicle
    if short != raw:
        vehicle = Vehicle.objects.filter(chassis_no=short).select_related('model', 'dealer').first()
        if vehicle:
            return vehicle
    # Only match longer stored values (legacy full VINs), never another short chassis
    return (
        Vehicle.objects.filter(chassis_no__endswith=short)
        .exclude(chassis_no=short)
        .filter(chassis_no__regex=r'.{8,}')
        .select_related('model', 'dealer')
        .order_by('-id')
        .first()
    )


def _as_text(value: Any) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    return str(value).strip()


def _normalize_header(value: Any) -> str:
    text = _as_text(value).lower()
    text = text.replace('\n', ' ')
    text = re.sub(r'\s+', ' ', text).strip()
    return text


def _as_date(value: Any):
    if value is None or value == '':
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = _as_text(value)
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y'):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _is_broken_customer(name: str) -> bool:
    upper = name.upper()
    return not name or upper.startswith('#REF') or upper in ('#N/A', '#VALUE!', 'NONE', 'NULL')


@dataclass
class ParsedRow:
    row_number: int
    status: str  # new | update | error
    chassis_no: str = ''
    customer: str = ''
    model_name: str = ''
    vehicle_type: str = ''
    lot: str = ''
    engine_no: str = ''
    arrival_date: str | None = None
    colour: str = ''
    bond_entry: str = ''
    receiving_status: str = ''
    model_id: int | None = None
    dealer_id: int | None = None
    vehicle_id: int | None = None
    new_customer: bool = False
    errors: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, data: dict) -> 'ParsedRow':
        return cls(**data)


def build_template_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Receiving'
    ws.append(TEMPLATE_HEADERS)
    return wb


def _map_headers(header_row) -> dict[str, list[int]]:
    mapping: dict[str, list[int]] = {key: [] for key in HEADER_ALIASES}
    for idx, cell in enumerate(header_row):
        normalized = _normalize_header(cell)
        if not normalized:
            continue
        for field_name, aliases in HEADER_ALIASES.items():
            if normalized in aliases:
                mapping[field_name].append(idx)
    return mapping


def _pick_chassis(row, chassis_indexes: list[int]) -> str:
    """Prefer a short (<=7 char) chassis column; else normalize the first chassis value."""
    if not chassis_indexes:
        return ''
    values = []
    for idx in chassis_indexes:
        if idx < len(row):
            values.append(_as_text(row[idx]))
    short = [v for v in values if v and len(v) <= 7]
    if short:
        return normalize_chassis(short[0])
    for value in values:
        if value:
            return normalize_chassis(value)
    return ''


def parse_receiving_workbook(file_obj) -> list[ParsedRow]:
    wb = load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb.active

    header_row = None
    header_row_number = 0
    mapping = {}
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        candidate = _map_headers(row)
        if candidate.get('chassis') or candidate.get('model'):
            header_row = row
            header_row_number = i
            mapping = candidate
            break

    if not header_row or not mapping.get('chassis'):
        raise ValueError('Could not find a header row with a CHASSIS NO column.')

    models_by_name = {m.name.upper(): m for m in VehicleModel.objects.all()}
    dealers_by_name = {d.name.upper(): d for d in Dealer.objects.all()}
    existing_chassis = {
        v.chassis_no: v
        for v in Vehicle.objects.select_related('model', 'dealer').all()
    }
    # Index longer (legacy full VIN) records by last-7 for update matching
    existing_by_suffix: dict[str, Vehicle] = {}
    for vehicle in existing_chassis.values():
        if len(vehicle.chassis_no) > 7:
            suffix = normalize_chassis(vehicle.chassis_no)
            existing_by_suffix.setdefault(suffix, vehicle)

    seen_in_file: dict[str, int] = {}
    results: list[ParsedRow] = []

    def cell(row, field_name, default=''):
        indexes = mapping.get(field_name) or []
        if not indexes:
            return default
        idx = indexes[0]
        if idx >= len(row):
            return default
        return row[idx]

    for row_number, row in enumerate(
        ws.iter_rows(min_row=header_row_number + 1, values_only=True),
        start=header_row_number + 1,
    ):
        if not any(row):
            continue

        parsed = ParsedRow(row_number=row_number, status='new')
        parsed.chassis_no = _pick_chassis(row, mapping.get('chassis') or [])
        parsed.customer = _as_text(cell(row, 'customer'))
        parsed.model_name = _as_text(cell(row, 'model'))
        parsed.vehicle_type = _as_text(cell(row, 'vehicle_type'))
        parsed.lot = _as_text(cell(row, 'lot'))
        parsed.engine_no = _as_text(cell(row, 'engine_no'))
        parsed.colour = _as_text(cell(row, 'colour'))
        parsed.bond_entry = _as_text(cell(row, 'bond_entry'))
        parsed.receiving_status = _as_text(cell(row, 'receiving_status'))

        arrival = _as_date(cell(row, 'arrival_date'))
        if arrival is not None:
            if arrival.year < 2000:
                parsed.errors.append(f'Invalid arrival date {arrival.isoformat()}')
            else:
                parsed.arrival_date = arrival.isoformat()

        if not parsed.chassis_no:
            parsed.errors.append('Missing chassis number')

        if parsed.chassis_no:
            if parsed.chassis_no in seen_in_file:
                parsed.errors.append(
                    f'Duplicate chassis in file (also on row {seen_in_file[parsed.chassis_no]})'
                )
            else:
                seen_in_file[parsed.chassis_no] = row_number

        if not parsed.model_name:
            parsed.errors.append('Missing model')
        else:
            model = models_by_name.get(parsed.model_name.upper())
            if not model:
                parsed.errors.append(f'Model "{parsed.model_name}" not found in master list')
            else:
                parsed.model_id = model.id

        if _is_broken_customer(parsed.customer):
            if parsed.customer:
                parsed.notes.append(f'Ignored invalid customer value "{parsed.customer}"')
            parsed.customer = ''
        elif parsed.customer:
            dealer = dealers_by_name.get(parsed.customer.upper())
            if dealer:
                parsed.dealer_id = dealer.id
            else:
                parsed.new_customer = True
                parsed.notes.append(f'New customer "{parsed.customer}" will be created')

        existing = existing_chassis.get(parsed.chassis_no) or existing_by_suffix.get(parsed.chassis_no)
        if existing and not parsed.errors:
            parsed.status = 'update'
            parsed.vehicle_id = existing.id
            parsed.notes.append('Existing unit — receiving details will be updated')
        elif parsed.errors:
            parsed.status = 'error'
        else:
            parsed.status = 'new'

        results.append(parsed)

    return results


def apply_receiving_rows(rows: list[ParsedRow], user=None) -> dict[str, int]:
    """Persist new/update rows. Returns counts."""
    created = updated = skipped = 0
    now = timezone.now()

    with transaction.atomic():
        dealers_by_name = {d.name.upper(): d for d in Dealer.objects.all()}
        models_by_id = {m.id: m for m in VehicleModel.objects.all()}

        for row in rows:
            if row.status == 'error':
                skipped += 1
                continue

            model = models_by_id.get(row.model_id) if row.model_id else None
            if not model:
                skipped += 1
                continue

            dealer = None
            if row.dealer_id:
                dealer = Dealer.objects.filter(pk=row.dealer_id).first()
            elif row.customer and row.new_customer:
                dealer = dealers_by_name.get(row.customer.upper())
                if not dealer:
                    dealer, _ = Dealer.objects.get_or_create(name=row.customer)
                    dealers_by_name[row.customer.upper()] = dealer

            arrival = None
            if row.arrival_date:
                arrival = datetime.date.fromisoformat(row.arrival_date)

            defaults = {
                'model': model,
                'dealer': dealer,
                'lot': row.lot or '',
                'vehicle_type': row.vehicle_type or '',
                'engine_no': row.engine_no or '',
                'colour': row.colour or '',
                'bond_entry': row.bond_entry or '',
                'arrival_date': arrival,
                'receiving_status': row.receiving_status or '',
                'received_by': user if getattr(user, 'is_authenticated', False) else None,
                'received_at': now,
            }

            if row.status == 'update' and row.vehicle_id:
                vehicle = Vehicle.objects.filter(pk=row.vehicle_id).first()
                if not vehicle:
                    skipped += 1
                    continue
                for key, value in defaults.items():
                    setattr(vehicle, key, value)
                # Keep existing chassis_no as-is for updates (may be full VIN)
                vehicle.save()
                updated += 1
            else:
                Vehicle.objects.update_or_create(
                    chassis_no=row.chassis_no,
                    defaults=defaults,
                )
                created += 1

    return {'created': created, 'updated': updated, 'skipped': skipped}
